import datetime
import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import plotly.express as px
import streamlit as st

# WeasyPrint check for PDF generation
try:
  from weasyprint import HTML

  WEASYPRINT_AVAILABLE = True
except Exception:
  WEASYPRINT_AVAILABLE = False

# Page Setup
st.set_page_config(
    page_title="E-Mojani Filtered Pending Report Generator & Dashboard",
    layout="wide",
)

# Custom Styling (CSS) for Streamlit UI
st.markdown(
    """
    <style>
    .metric-card {
        background-color: #f8f9fa;
        border-radius: 12px;
        padding: 18px;
        border-left: 6px solid #2b9348;
        box-shadow: 0 4px 6px rgba(0,0,0,0.05);
        text-align: center;
    }
    .metric-title { font-size: 14px; color: #555; font-weight: 600; }
    .metric-value { font-size: 26px; color: #1b4332; font-weight: bold; margin-top: 5px; }
    </style>
""",
    unsafe_allow_html=True,
)

st.title("📊 भूमि अभिलेख विभाग - अमरावती | प्रलंबित मोजणी अहवाल व डॅशबोर्ड")


# EXCEL HELPER FOR INDIVIDUAL SHEETS (WITH TITLE HEADER)
def convert_df_to_excel(df, report_title="प्रलंबित मोजणी अहवाल"):
  output = io.BytesIO()
  current_date_str = datetime.date.today().strftime("%d/%m/%Y")

  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Report"

  font_dept = Font(name="Calibri", size=14, bold=True, color="1F4E78")
  font_sub = Font(name="Calibri", size=11, bold=True, color="000000")
  font_th = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
  fill_th = PatternFill(
      start_color="1F4E78", end_color="1F4E78", fill_type="solid"
  )

  thin_border = Border(
      left=Side(style="thin", color="CCCCCC"),
      right=Side(style="thin", color="CCCCCC"),
      top=Side(style="thin", color="CCCCCC"),
      bottom=Side(style="thin", color="CCCCCC"),
  )

  is_multi = isinstance(df.index, pd.MultiIndex)
  num_idx_cols = len(df.index.names) if is_multi else 0
  total_cols = num_idx_cols + len(df.columns)

  # Row 1: Department Title
  ws.merge_cells(
      start_row=1, start_column=1, end_row=1, end_column=max(total_cols, 2)
  )
  cell_dept = ws.cell(row=1, column=1, value="भूमि अभिलेख विभाग - अमरावती")
  cell_dept.font = font_dept
  cell_dept.alignment = Alignment(horizontal="center", vertical="center")

  # Row 2: Subtitle with Date
  ws.merge_cells(
      start_row=2, start_column=1, end_row=2, end_column=max(total_cols, 2)
  )
  cell_sub = ws.cell(
      row=2, column=1, value=f"{report_title} ( दिनांक {current_date_str} )"
  )
  cell_sub.font = font_sub
  cell_sub.alignment = Alignment(horizontal="center", vertical="center")

  # Table Headers
  start_row = 4
  col_offset = 1

  if is_multi:
    for i, idx_name in enumerate(df.index.names):
      cell = ws.cell(row=start_row, column=i + 1, value=idx_name or "")
      cell.font, cell.fill = font_th, fill_th
      cell.alignment = Alignment(horizontal="center", vertical="center")
    col_offset = num_idx_cols + 1

  for j, col_name in enumerate(df.columns):
    cell = ws.cell(row=start_row, column=col_offset + j, value=str(col_name))
    cell.font, cell.fill = font_th, fill_th
    cell.alignment = Alignment(horizontal="center", vertical="center")

  # Rows Data
  current_r = start_row + 1
  for idx_val, row in df.iterrows():
    if is_multi:
      for k, val in enumerate(idx_val):
        cell = ws.cell(row=current_r, column=k + 1, value=str(val))
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

    for j, col_name in enumerate(df.columns):
      val = row[col_name]
      val_clean = (
          ""
          if (
              pd.isna(val)
              or val == 0
              or str(val).strip() in ["", "0", "nan", "0.0"]
          )
          else (int(val) if isinstance(val, (int, float)) else val)
      )
      cell = ws.cell(row=current_r, column=col_offset + j, value=val_clean)
      cell.alignment = Alignment(horizontal="center", vertical="center")
      cell.border = thin_border

    current_r += 1

  # Auto Column Width
  for col in ws.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

  wb.save(output)
  return output.getvalue()


# MASTER EXCEL HELPER: COMBINES ALL REPORTS INTO SINGLE EXCEL
def convert_all_reports_to_single_excel(df_dict):
  output = io.BytesIO()
  current_date_str = datetime.date.today().strftime("%d/%m/%Y")
  wb = openpyxl.Workbook()
  wb.remove(wb.active)

  font_dept = Font(name="Calibri", size=14, bold=True, color="1F4E78")
  font_sub = Font(name="Calibri", size=11, bold=True, color="000000")
  font_th = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
  fill_th = PatternFill(
      start_color="1F4E78", end_color="1F4E78", fill_type="solid"
  )

  thin_border = Border(
      left=Side(style="thin", color="CCCCCC"),
      right=Side(style="thin", color="CCCCCC"),
      top=Side(style="thin", color="CCCCCC"),
      bottom=Side(style="thin", color="CCCCCC"),
  )

  for sheet_name, (df_data, report_title) in df_dict.items():
    ws = wb.create_sheet(title=sheet_name[:31])

    is_multi = isinstance(df_data.index, pd.MultiIndex)
    num_idx_cols = len(df_data.index.names) if is_multi else 0
    total_cols = num_idx_cols + len(df_data.columns)

    # Title Headers
    ws.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=max(total_cols, 2)
    )
    c1 = ws.cell(row=1, column=1, value="भूमि अभिलेख विभाग - अमरावती")
    c1.font, c1.alignment = font_dept, Alignment(
        horizontal="center", vertical="center"
    )

    ws.merge_cells(
        start_row=2, start_column=1, end_row=2, end_column=max(total_cols, 2)
    )
    c2 = ws.cell(
        row=2, column=1, value=f"{report_title} ( दिनांक {current_date_str} )"
    )
    c2.font, c2.alignment = font_sub, Alignment(
        horizontal="center", vertical="center"
    )

    start_row = 4
    col_offset = 1

    if is_multi:
      for i, idx_name in enumerate(df_data.index.names):
        cell = ws.cell(row=start_row, column=i + 1, value=idx_name or "")
        cell.font, cell.fill = font_th, fill_th
        cell.alignment = Alignment(horizontal="center", vertical="center")
      col_offset = num_idx_cols + 1

    for j, col_name in enumerate(df_data.columns):
      cell = ws.cell(row=start_row, column=col_offset + j, value=str(col_name))
      cell.font, cell.fill = font_th, fill_th
      cell.alignment = Alignment(horizontal="center", vertical="center")

    current_r = start_row + 1
    for idx_val, row in df_data.iterrows():
      if is_multi:
        for k, val in enumerate(idx_val):
          cell = ws.cell(row=current_r, column=k + 1, value=str(val))
          cell.alignment = Alignment(horizontal="left", vertical="center")
          cell.border = thin_border

      for j, col_name in enumerate(df_data.columns):
        val = row[col_name]
        val_clean = (
            ""
            if (
                pd.isna(val)
                or val == 0
                or str(val).strip() in ["", "0", "nan", "0.0"]
            )
            else (int(val) if isinstance(val, (int, float)) else val)
        )
        cell = ws.cell(row=current_r, column=col_offset + j, value=val_clean)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

      current_r += 1

    for col in ws.columns:
      max_len = max(len(str(cell.value or "")) for cell in col)
      col_letter = openpyxl.utils.get_column_letter(col[0].column)
      ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

  wb.save(output)
  return output.getvalue()


# PROFESSIONAL & CLEAN PDF GENERATOR FOR ALL REPORTS
def generate_report_pdf_html(df, subtitle_text=""):
  headers = list(df.columns)
  current_date_str = datetime.date.today().strftime("%d/%m/%Y")
  is_multi_index = isinstance(df.index, pd.MultiIndex)

  html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            @page {{ size: A4 landscape; margin: 8mm; }}
            body {{ font-family: 'Helvetica Neue', 'Arial', sans-serif; margin: 0; padding: 0; background-color: #ffffff; color: #111; }}
            .header-container {{ text-align: center; margin-bottom: 8px; padding-bottom: 4px; border-bottom: 2px solid #1a365d; }}
            .dept-title {{ font-size: 15px; font-weight: bold; color: #1a365d; text-transform: uppercase; margin-bottom: 2px; }}
            .report-title {{ font-size: 12px; font-weight: 600; color: #2b6cb0; }}
            table {{ width: 100%; border-collapse: collapse; border: 1px solid #4a5568; margin-top: 4px; font-size: 8.5px; table-layout: auto; }}
            thead {{ display: table-header-group; }}
            tr {{ page-break-inside: avoid; }}
            th {{ background-color: #1a365d; color: #ffffff; font-weight: bold; font-size: 8.5px; padding: 4px 2px; border: 1px solid #2d3748; text-align: center; line-height: 1.1; }}
            td {{ padding: 3px 2px; font-size: 8px; border: 1px solid #cbd5e0; text-align: center; color: #2d3748; line-height: 1.1; }}
            tbody tr:nth-child(even) {{ background-color: #f7fafc; }}
            .text-col {{ text-align: left !important; padding-left: 4px !important; font-weight: 500; }}
            .total-row td {{ font-weight: bold !important; background-color: #edf2f7 !important; color: #1a365d !important; border-top: 1.5px solid #1a365d !important; }}
        </style>
    </head>
    <body>
        <div class="header-container">
            <div class="dept-title">भूमि अभिलेख विभाग - अमरावती</div>
            <div class="report-title">{subtitle_text} ( दिनांक {current_date_str} )</div>
        </div>
        <table>
            <thead>
                <tr>
    """

  if is_multi_index:
    for idx_name in df.index.names:
      html += f"<th>{idx_name}</th>"

  for h in headers:
    html += f"<th>{h}</th>"

  html += """
                </tr>
            </thead>
            <tbody>
    """

  for idx, row in df.iterrows():
    idx_str = str(idx)
    is_total = (
        "Total" in idx_str or "एकूण" in idx_str or "Grand Total" in idx_str
    )
    row_class = "total-row" if is_total else ""

    html += f'<tr class="{row_class}">'

    if is_multi_index:
      for idx_item in idx:
        html += f'<td class="text-col">{idx_item}</td>'

    for col_name in headers:
      val = row[col_name]
      val_str = (
          ""
          if (
              pd.isna(val)
              or val == 0
              or str(val).strip() in ["", "0", "nan", "0.0"]
          )
          else str(int(val) if isinstance(val, (int, float)) else val)
      )
      html += f"<td>{val_str}</td>"

    html += "</tr>"

  html += """
            </tbody>
        </table>
    </body>
    </html>
    """

  if WEASYPRINT_AVAILABLE:
    try:
      return HTML(string=html).write_pdf()
    except Exception:
      return html.encode("utf-8")
  else:
    return html.encode("utf-8")


# Sidebar File Upload
st.sidebar.header("⚙️ File Upload & Options")
uploaded_file = st.sidebar.file_uploader(
    "📁 Main E-Mojani Raw Excel (.xlsx)", type=["xlsx"], key="single_file"
)


def parse_excel_date(val):
  if pd.isna(val) or str(val).strip() in ["", "nan", "NaT"]:
    return pd.NaT
  try:
    val_num = float(val)
    return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=val_num)
  except Exception:
    return pd.to_datetime(val, dayfirst=True, errors="coerce")


if uploaded_file is not None:
  df_raw = pd.read_excel(uploaded_file)
  if "मोजणी तारीख" not in df_raw.columns:
    df_raw = pd.read_excel(uploaded_file, header=1)

  if "तालुका" in df_raw.columns:
    df_raw["तालुका"] = df_raw["तालुका"].astype(str).str.strip()
    df_raw = df_raw[
        df_raw["तालुका"].notna()
        & (~df_raw["तालुका"].isin(["", "nan", "None", "0"]))
    ]

  col_mojni_date = (
      "मोजणी तारीख"
      if "मोजणी तारीख" in df_raw.columns
      else df_raw.columns[8]
      if len(df_raw.columns) >= 9
      else None
  )
  col_status = (
      "स्थिती"
      if "स्थिती" in df_raw.columns
      else df_raw.columns[10]
      if len(df_raw.columns) >= 11
      else None
  )

  df_raw["mojni_date_parsed"] = df_raw[col_mojni_date].apply(parse_excel_date)

  col_k_prat_date = "क प्रत निकाली दिनाक"
  if col_k_prat_date in df_raw.columns:
    df_raw["k_prat_date_parsed"] = df_raw[col_k_prat_date].apply(
        parse_excel_date
    )
  else:
    df_raw["k_prat_date_parsed"] = df_raw["mojni_date_parsed"]

  # Parse Application Accept Date
  col_app_accept_date_raw = next(
      (
          c
          for c in ["अर्ज स्वीकृत दिनांक", "स्वीकृत दिनांक"]
          if c in df_raw.columns
      ),
      None,
  )
  if col_app_accept_date_raw:
    df_raw["app_accept_date_parsed"] = df_raw[col_app_accept_date_raw].apply(
        parse_excel_date
    )
  else:
    df_raw["app_accept_date_parsed"] = df_raw["mojni_date_parsed"]

  min_excel_date = df_raw["mojni_date_parsed"].min()
  max_excel_date = df_raw["mojni_date_parsed"].max()

  min_excel_date = (
      datetime.date(2023, 1, 1)
      if pd.isna(min_excel_date)
      else min_excel_date.date()
  )
  max_excel_date = (
      datetime.date.today()
      if pd.isna(max_excel_date)
      else max_excel_date.date()
  )
  today_file_str = datetime.date.today().strftime("%d-%m-%Y")

  st.sidebar.markdown("---")
  st.sidebar.subheader("📅 Report 1 कालावधी (Date Range)")
  col_d1, col_d2 = st.sidebar.columns(2)
  start_date = col_d1.date_input(
      "पासून (From)", value=min_excel_date, key="r1_start"
  )
  end_date = col_d2.date_input(
      "पर्यंत (To)", value=datetime.date.today(), key="r1_end"
  )

  st.sidebar.markdown("---")
  st.sidebar.subheader("📌 स्थिती (Status) Filter")
  all_statuses = df_raw[col_status].dropna().unique().tolist()
  completed_defaults = [
      "क प्रत",
      "विनाकार्यवाही",
      "प्रस्तावित बिगरशेती/गुंठेवारी मोजणी पूर्ण",
  ]
  default_selected = [s for s in all_statuses if s not in completed_defaults]

  selected_statuses = st.sidebar.multiselect(
      "Select Status:", options=all_statuses, default=default_selected
  )

  with st.spinner("डेटा प्रोसेस होत आहे... कृपया वाट पहा..."):
    tab1, tab2, tab3, tab4, tab5, tab6, tab7, tab8 = st.tabs([
        "📊 Report 1 (तालुका व दिवसनिहाय)",
        "📋 Report 2 (टप्पा व अधिकारी निहाय)",
        "📈 Executive Dashboard",
        "📋 Report 3 (भूकरमापक निहाय)",
        "📅 Report 4 (मोजणी होणेवर शिल्लक)",
        "🚀 Report 5 (मागील ३ महिने क प्रत)",
        "⚠️ Report 6 (भूकरमापक १०+ प्रलंबित)",
        "📑 Report 7 (वैयक्तिक यादी)",
    ])

    from_str = start_date.strftime("%d/%m/%Y")
    to_str = end_date.strftime("%d/%m/%Y")

    # --- REPORT 1 GENERATION ---
    df1 = df_raw[df_raw[col_status].isin(selected_statuses)].copy()
    start_dt, end_dt = pd.to_datetime(start_date), pd.to_datetime(end_date)
    df1 = df1[
        (df1["mojni_date_parsed"] >= start_dt)
        & (df1["mojni_date_parsed"] <= end_dt)
    ].copy()
    df1["Pending Days"] = (end_dt - df1["mojni_date_parsed"]).dt.days

    def day_bucket(days):
      if pd.isna(days):
        return "तारीख उपलब्ध नाही"
      if days <= 15:
        return "15 दिवस वर"
      elif days <= 30:
        return "30 दिवस वर"
      elif days <= 60:
        return "60 दिवस वर"
      elif days <= 90:
        return "90 दिवस वर"
      else:
        return "90 दिवसांपेक्षा जास्त"

    df1["दिवस वर"] = df1["Pending Days"].apply(day_bucket)
    bucket_order = [
        "15 दिवस वर",
        "30 दिवस वर",
        "60 दिवस वर",
        "90 दिवस वर",
        "90 दिवसांपेक्षा जास्त",
    ]
    pivot_df = pd.crosstab(
        df1["तालुका"], df1["दिवस वर"], margins=True, margins_name="एकूण"
    )
    existing_cols = [c for c in bucket_order if c in pivot_df.columns]
    if "तारीख उपलब्ध नाही" in pivot_df.columns:
      existing_cols.append("तारीख उपलब्ध नाही")
    if "एकूण" in pivot_df.columns:
      existing_cols.append("एकूण")
    pivot_df = pivot_df[existing_cols]

    # --- REPORT 2 GENERATION ---
    r2_start_dt, r2_end_dt = pd.to_datetime(min_excel_date), pd.to_datetime(
        datetime.date.today()
    )
    today_dt = pd.to_datetime(datetime.date.today())
    df_r2_filtered = df_raw[
        (df_raw["mojni_date_parsed"] >= r2_start_dt)
        & (df_raw["mojni_date_parsed"] <= r2_end_dt)
    ].copy()
    col_yn, col_mojni_type = (
        "क्षेत्र अभिलेखाशी मेळात आहे का?",
        "मोजणीचा प्रकार(Mojni Type)",
    )
    taluka_list = sorted(df_raw["तालुका"].dropna().unique().tolist())
    report2_data = []

    for tal in taluka_list:
      df_t = df_r2_filtered[df_r2_filtered["तालुका"] == tal]
      df_taluka_all = df_raw[df_raw["तालुका"] == tal]
      chanani = len(df_t[df_t[col_status] == "छाननी लिपिक यांनी तपासले"])
      shirastedar = len(
          df_t[
              df_t[col_status]
              == "शिरस्तेदार/मुख्यालय सहाय्यक यांनी तपासले"
          ]
      )
      up_bhoo = len(
          df_t[
              df_t[col_status] == "ऊप.अ. भू. अ/ न .भू अ यांच्या मान्यतेवर"
          ]
      )
      off_total = chanani + shirastedar + up_bhoo

      yes_no = (
          len(
              df_t[
                  df_t[col_yn].notna()
                  & (df_t[col_yn].astype(str).str.strip() != "")
              ]
          )
          if col_yn in df_t.columns
          else 0
      )
      if col_mojni_type in df_t.columns:
        haddi = len(
            df_t[
                (df_t[col_mojni_type] == "ह्द्दकायम")
                & (df_t[col_status] == "मोजणीची माहिती")
            ]
        )
        jama = len(
            df_t[
                (df_t[col_status] == "मोजणीची माहिती")
                & (df_t[col_mojni_type] != "ह्द्दकायम")
            ]
        )
      else:
        haddi = len(df_t[df_t[col_status] == "मोजणीची माहिती"])
        jama = 0

      shillak = len(
          df_taluka_all[
              (df_taluka_all["mojni_date_parsed"] > today_dt)
              & (df_taluka_all[col_status] == "मोजणीची माहिती")
          ]
      )
      stage_total = yes_no + jama + haddi + shillak

      report2_data.append({
          "तालुका": tal,
          "Yes/No": yes_no,
          "जमा करणेवर": jama,
          "हददी दाखविणेवर": haddi,
          "शिल्लक प्रकरणे": shillak,
          "Stage Total": stage_total,
          "छाननी लिपीक": chanani,
          "शिरस्तेदार/मुख्यालय सहाय्यक": shirastedar,
          "उप अ भू अ/ भू अ": up_bhoo,
          "Officer Total": off_total,
      })

    df_rep2 = pd.DataFrame(report2_data)
    total_row_r2 = {
        "तालुका": "एकूण",
        "Yes/No": df_rep2["Yes/No"].sum(),
        "जमा करणेवर": df_rep2["जमा करणेवर"].sum(),
        "हददी दाखविणेवर": df_rep2["हददी दाखविणेवर"].sum(),
        "शिल्लक प्रकरणे": df_rep2["शिल्लक प्रकरणे"].sum(),
        "Stage Total": df_rep2["Stage Total"].sum(),
        "छाननी लिपीक": df_rep2["छाननी लिपीक"].sum(),
        "शिरस्तेदार/मुख्यालय सहाय्यक": df_rep2[
            "शिरस्तेदार/मुख्यालय सहाय्यक"
        ].sum(),
        "उप अ भू अ/ भू अ": df_rep2["उप अ भू अ/ भू अ"].sum(),
        "Officer Total": df_rep2["Officer Total"].sum(),
    }
    df_rep2 = pd.concat(
        [df_rep2, pd.DataFrame([total_row_r2])], ignore_index=True
    ).set_index("तालुका")

    # --- REPORT 3 GENERATION ---
    surveyor_col_r3 = next(
        (
            col
            for col in ["भूकरमापक", "कर्मचारी/अधिकारी चे नाव"]
            if col in df_raw.columns
        ),
        None,
    )
    pivot_r3 = pd.DataFrame()
    if surveyor_col_r3:
      df_r3_clean = df_raw[
          df_raw["तालुका"].notna() & df_raw[surveyor_col_r3].notna()
      ].copy()
      df_r3_pending = df_r3_clean[
          df_r3_clean[col_status].isin(selected_statuses)
      ].copy()
      curr_today = pd.to_datetime(datetime.date.today())
      df_r3_pending = df_r3_pending[
          df_r3_pending["mojni_date_parsed"].notna()
          & (df_r3_pending["mojni_date_parsed"] <= curr_today)
      ].copy()

      if not df_r3_pending.empty:
        df_r3_pending["R3_Pending_Days"] = (
            curr_today - df_r3_pending["mojni_date_parsed"]
        ).dt.days

        def r3_ageing_bucket(days):
          if pd.isna(days):
            return "(blank)"
          elif days <= 15:
            return "0–15 दिवस वर"
          elif days <= 30:
            return "16–30 दिवस वर"
          elif days <= 60:
            return "31–60 दिवस वर"
          elif days <= 90:
            return "61–90 दिवस वर"
          else:
            return "90 दिवसांपेक्षा जास्त"

        df_r3_pending["Ageing_Bucket"] = df_r3_pending["R3_Pending_Days"].apply(
            r3_ageing_bucket
        )
        pivot_r3 = pd.pivot_table(
            df_r3_pending,
            index=["तालुका", surveyor_col_r3],
            columns="Ageing_Bucket",
            values=(
                "अर्ज क्र.(Application No)"
                if "अर्ज क्र.(Application No)" in df_r3_pending.columns
                else df_r3_pending.columns[0]
            ),
            aggfunc="count",
            fill_value=0,
            margins=True,
            margins_name="Grand Total",
        )
        r3_bucket_order = [
            "0–15 दिवस वर",
            "16–30 दिवस वर",
            "31–60 दिवस वर",
            "61–90 दिवस वर",
            "90 दिवसांपेक्षा जास्त",
            "(blank)",
            "Grand Total",
        ]
        pivot_r3 = pivot_r3[[c for c in r3_bucket_order if c in pivot_r3.columns]]

    # --- REPORT 4 GENERATION ---
    today_ref = pd.to_datetime(datetime.date.today())
    df_r4 = df_raw[
        (df_raw["mojni_date_parsed"].notna())
        & (df_raw["mojni_date_parsed"] > today_ref)
        & (df_raw[col_status].astype(str).str.contains("मोजणीची माहिती"))
    ].copy()
    final_r4_df = pd.DataFrame()
    if not df_r4.empty:
      df_r4["Month_Year"] = df_r4["mojni_date_parsed"].dt.strftime("%b")
      df_r4["Year_Month_Sort"] = df_r4["mojni_date_parsed"].dt.strftime(
          "%Y-%m"
      )
      sorted_months_df = (
          df_r4[["Year_Month_Sort", "Month_Year"]]
          .drop_duplicates()
          .sort_values("Year_Month_Sort")
      )
      sorted_month_names = sorted_months_df["Month_Year"].tolist()

      pivot_r4 = pd.pivot_table(
          df_r4,
          index="तालुका",
          columns="Month_Year",
          values=df_r4.columns[0],
          aggfunc="count",
          fill_value=0,
      )
      pivot_r4 = pivot_r4[
          [m for m in sorted_month_names if m in pivot_r4.columns]
      ]

      standard_talukas = [
          "अंजनगाव",
          "अचलपूर",
          "अमरावती",
          "चांदुर बा.",
          "चांदुर रे.",
          "चिखलदरा",
          "तिवसा",
          "दर्यापूर",
          "धामणगाव",
          "धारणी",
          "नांदगाव",
          "भातकुली",
          "मोर्शी",
          "वरुड",
      ]
      pivot_r4 = pivot_r4.rename(
          index={"अचलपुर": "अचलपूर", "दर्यापुर": "दर्यापूर"}
      )
      for t in pivot_r4.index.tolist():
        if t not in standard_talukas:
          standard_talukas.append(t)

      pivot_r4 = pivot_r4.reindex(standard_talukas, fill_value=0)
      pivot_r4["Grand Total"] = pivot_r4.sum(axis=1)

      district_total = pivot_r4.sum(axis=0)
      df_district_row = pd.DataFrame(
          [district_total], index=["अमरावती (जिल्हा एकूण)"]
      )
      final_r4_df = pd.concat([df_district_row, pivot_r4])

    # --- REPORT 5 GENERATION ---
    today_ref_r5 = pd.to_datetime(datetime.date.today())
    three_months_ago_r5 = today_ref_r5 - pd.Timedelta(days=90)
    df_r5 = df_raw[
        (df_raw["k_prat_date_parsed"].notna())
        & (df_raw["k_prat_date_parsed"] >= three_months_ago_r5)
        & (df_raw["k_prat_date_parsed"] <= today_ref_r5)
        & (df_raw[col_status].astype(str).str.strip() == "क प्रत")
    ].copy()
    pivot_r5 = pd.DataFrame()
    if not df_r5.empty:
      df_r5["Month_Year"] = df_r5["k_prat_date_parsed"].dt.strftime("%b")
      df_r5["Year_Month_Sort"] = df_r5["k_prat_date_parsed"].dt.strftime(
          "%Y-%m"
      )
      sorted_months_r5 = (
          df_r5[["Year_Month_Sort", "Month_Year"]]
          .drop_duplicates()
          .sort_values("Year_Month_Sort")
      )
      sorted_month_names_r5 = sorted_months_r5["Month_Year"].tolist()

      pivot_r5 = pd.pivot_table(
          df_r5,
          index="तालुका",
          columns="Month_Year",
          values=df_r5.columns[0],
          aggfunc="count",
          fill_value=0,
      )
      pivot_r5 = pivot_r5[
          [m for m in sorted_month_names_r5 if m in pivot_r5.columns]
      ]

      standard_talukas = [
          "अंजनगाव",
          "अचलपूर",
          "अमरावती",
          "चांदुर बा.",
          "चांदुर रे.",
          "चिखलदरा",
          "तिवसा",
          "दर्यापूर",
          "धामणगाव",
          "धारणी",
          "नांदगाव",
          "भातकुली",
          "मोर्शी",
          "वरुड",
      ]
      pivot_r5 = pivot_r5.rename(
          index={"अचलपुर": "अचलपूर", "दर्यापुर": "दर्यापूर"}
      )
      for t in pivot_r5.index.tolist():
        if t not in standard_talukas:
          standard_talukas.append(t)

      pivot_r5 = pivot_r5.reindex(standard_talukas, fill_value=0)
      pivot_r5["Grand Total"] = pivot_r5.sum(axis=1)

      total_row_r5 = pivot_r5.sum(axis=0)
      pivot_r5.loc["Total"] = total_row_r5

    # --- REPORT 6 GENERATION ---
    df_r6 = pd.DataFrame()
    if not pivot_r3.empty:
      df_r6_temp = pivot_r3.reset_index()
      df_r6_temp = df_r6_temp[df_r6_temp["तालुका"] != "Grand Total"].copy()
      if "Grand Total" in df_r6_temp.columns:
        df_r6_temp = df_r6_temp[df_r6_temp["Grand Total"] >= 10].copy()
        df_r6 = df_r6_temp[["तालुका", surveyor_col_r3, "Grand Total"]].set_index(
            ["तालुका", surveyor_col_r3]
        )

    # --- REPORT 7 GENERATION (4 SPECIFIC STATUSES & ACC DATE TO MOJNI DATE DAYS) ---
    r7_target_statuses = [
        "ऊप.अ. भू. अ/ न .भू अ यांच्या मान्यतेवर",
        "छाननी लिपिक यांनी तपासले",
        "शिरस्तेदार/मुख्यालय सहाय्यक यांनी तपासले",
        "मोजणीची माहिती",
    ]

    df_r7 = df_raw[
        df_raw[col_status].astype(str).str.strip().isin(r7_target_statuses)
    ].copy()

    col_app_no = next(
        (c for c in ["अर्ज क्र.(Application No)", "अर्ज क्र."] if c in df_r7.columns),
        df_r7.columns[0],
    )
    col_dist = next(
        (c for c in ["जिल्हा", "District"] if c in df_r7.columns), None
    )
    col_village = next(
        (c for c in ["गाव", "Village"] if c in df_r7.columns), None
    )

    # Calculate Days: Mojni Date - Application Accept Date
    df_r7["दिवस"] = (
        df_r7["mojni_date_parsed"] - df_r7["app_accept_date_parsed"]
    ).dt.days

    def r7_kalavadhi(days):
      if pd.isna(days) or days < 0:
        return "-"
      elif days <= 15:
        return "15-30" if days > 15 else "0-15"
      elif days <= 30:
        return "15-30"
      elif days <= 60:
        return "30-60"
      elif days <= 90:
        return "60-90"
      elif days <= 180:
        return "90-180"
      else:
        return "180+"

    df_r7["कालावधी"] = df_r7["दिवस"].apply(r7_kalavadhi)

    # Format dates to string
    df_r7["मोजणी तारीख प्रदर्शित"] = df_r7["mojni_date_parsed"].dt.strftime(
        "%d-%m-%Y"
    )
    df_r7["अर्ज स्वीकृत दिनांक प्रदर्शित"] = df_r7[
        "app_accept_date_parsed"
    ].dt.strftime("%d-%m-%Y")

    r7_cols_map = {}
    r7_cols_map[col_app_no] = "अर्ज क्र.(Application No)"
    if col_dist:
      r7_cols_map[col_dist] = "जिल्हा"
    else:
      df_r7["जिल्हा"] = "अमरावती"
      r7_cols_map["जिल्हा"] = "जिल्हा"
    r7_cols_map["तालुका"] = "तालुका"
    if col_village:
      r7_cols_map[col_village] = "गाव"
    if surveyor_col_r3:
      r7_cols_map[surveyor_col_r3] = "भूकरमापक"
    r7_cols_map[col_status] = "स्थिती"
    r7_cols_map["मोजणी तारीख प्रदर्शित"] = "मोजणी तारीख"
    r7_cols_map["अर्ज स्वीकृत दिनांक प्रदर्शित"] = "अर्ज स्वीकृत दिनांक"
    r7_cols_map["दिवस"] = "दिवस"
    r7_cols_map["कालावधी"] = "कालावधी"

    df_r7_display = df_r7[list(r7_cols_map.keys())].rename(columns=r7_cols_map)

    # SIDEBAR: MASTER ALL-IN-ONE EXCEL EXPORT BUTTON
    st.sidebar.markdown("---")
    st.sidebar.subheader("📦 Master Single Excel Download")

    all_reports_dict = {
        "Report_1": (
            pivot_df,
            "तालुका व दिवसनिहाय प्रलंबित मोजणी प्रकरणे अहवाल",
        ),
        "Report_2": (
            df_rep2,
            "टप्पा व अधिकारी निहाय प्रलंबित मोजणी प्रकरणे अहवाल",
        ),
        "Report_3": (
            pivot_r3,
            "भूकरमापक निहाय व दिवसनिहाय प्रलंबित मोजणी प्रकरणे अहवाल",
        ),
        "Report_4": (final_r4_df, "मोजणी होणेवर शिल्लक प्रकरणांचा अहवाल"),
        "Report_5": (pivot_r5, "मागील ३ महिने 'क प्रत' प्रोग्रेस अहवाल"),
        "Report_6": (
            df_r6,
            "भूकरमापक प्रलंबित मोजणी १० पेक्षा जास्त प्रकरणे अहवाल",
        ),
        "Report_7": (
            df_r7_display,
            "प्रलंबित प्रकरणांची वैयक्तिक यादी (जिल्हा ➔ तालुका ➔ भूकरमापक)"
            " अहवाल",
        ),
    }

    master_excel_bytes = convert_all_reports_to_single_excel(all_reports_dict)

    st.sidebar.download_button(
        label="📊 Download All 7 Reports in 1 Excel",
        data=master_excel_bytes,
        file_name=f"All_Mojni_Reports_{today_file_str}.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )

    # TABS DISPLAY LOGIC
    with tab1:
      st.success(f"✅ Total Filtered Cases: **{len(df1)}**")
      st.subheader("📋 तालुका व दिवसनिहाय प्रलंबित प्रकरणे Table")
      st.dataframe(pivot_df, use_container_width=True)

      btn_col1, btn_col2 = st.columns(2)
      with btn_col1:
        excel_r1 = convert_df_to_excel(
            pivot_df,
            report_title="तालुका व दिवसनिहाय प्रलंबित मोजणी प्रकरणे अहवाल",
        )
        st.download_button(
            label="📥 Download Report 1 (Excel)",
            data=excel_r1,
            file_name=f"Report_1_Talukani_hay_{today_file_str}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
      with btn_col2:
        pdf_r1 = generate_report_pdf_html(
            pivot_df,
            subtitle_text="तालुका व दिवसनिहाय प्रलंबित मोजणी प्रकरणे अहवाल",
        )
        st.download_button(
            label="📄 Download Report 1 (PDF / Print)",
            data=pdf_r1,
            file_name=f"Report_1_Talukani_hay_{today_file_str}.pdf",
            mime="application/pdf" if WEASYPRINT_AVAILABLE else "text/html",
        )

    with tab2:
      st.subheader("📋 Report 2 - टप्पा व अधिकारी निहाय अहवाल")
      st.dataframe(df_rep2, use_container_width=True)

      btn_r2_1, btn_r2_2 = st.columns(2)
      with btn_r2_1:
        excel_r2 = convert_df_to_excel(
            df_rep2,
            report_title="टप्पा व अधिकारी निहाय प्रलंबित मोजणी प्रकरणे अहवाल",
        )
        st.download_button(
            label="📥 Download Report 2 (Excel)",
            data=excel_r2,
            file_name=f"Report_2_Stage_Officer_{today_file_str}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
      with btn_r2_2:
        pdf_r2 = generate_report_pdf_html(
            df_rep2,
            subtitle_text="टप्पा व अधिकारी निहाय प्रलंबित मोजणी प्रकरणे अहवाल",
        )
        st.download_button(
            label="📄 Download Report 2 (PDF / Print)",
            data=pdf_r2,
            file_name=f"Report_2_Stage_Officer_{today_file_str}.pdf",
            mime="application/pdf" if WEASYPRINT_AVAILABLE else "text/html",
        )

    with tab3:
      st.subheader("📈 Executive Dashboard Overview")
      col_m1, col_m2, col_m3 = st.columns(3)
      with col_m1:
        st.markdown(
            f'<div class="metric-card"><div class="metric-title">एकूण अपलोड'
            f' प्रकरणे</div><div class="metric-value">{len(df_raw)}</div></div>',
            unsafe_allow_html=True,
        )
      with col_m2:
        st.markdown(
            '<div class="metric-card"><div class="metric-title">प्रलंबित'
            ' प्रकरणे (Filtered)</div><div'
            f' class="metric-value">{len(df1)}</div></div>',
            unsafe_allow_html=True,
        )
      with col_m3:
        today_ref_dash = pd.to_datetime(datetime.date.today())
        r4_count_dash = len(
            df_raw[
                (df_raw["mojni_date_parsed"].notna())
                & (df_raw["mojni_date_parsed"] > today_ref_dash)
                & (
                    df_raw[col_status]
                    .astype(str)
                    .str.contains("मोजणीची माहिती")
                )
            ]
        )
        st.markdown(
            '<div class="metric-card"><div class="metric-title">मोजणी होणेवर'
            f' शिल्लक</div><div class="metric-value">{r4_count_dash}</div></div>',
            unsafe_allow_html=True,
        )

      st.markdown("---")
      if "तालुका" in df_raw.columns:
        st.subheader("📊 तालुका निहाय प्रलंबित प्रकरणे चार्ट")
        tal_chart_df = df1["तालुका"].value_counts().reset_index()
        tal_chart_df.columns = ["तालुका", "प्रकरणे"]
        fig_tal = px.bar(
            tal_chart_df,
            x="तालुका",
            y="प्रकरणे",
            color="प्रकरणे",
            text_auto=True,
            color_continuous_scale="Viridis",
        )
        st.plotly_chart(fig_tal, use_container_width=True)

    with tab4:
      st.markdown("## 📋 रिपोर्ट ३: भूकरमापक निहाय व दिवसनिहाय")
      if not pivot_r3.empty:
        st.dataframe(pivot_r3, use_container_width=True)

        btn_r3_1, btn_r3_2 = st.columns(2)
        with btn_r3_1:
          excel_r3 = convert_df_to_excel(
              pivot_r3,
              report_title=(
                  "भूकरमापक निहाय व दिवसनिहाय प्रलंबित मोजणी प्रकरणे अहवाल"
              ),
          )
          st.download_button(
              label="📥 Download Report 3 (Excel)",
              data=excel_r3,
              file_name=f"Report_3_Bhukarmapak_{today_file_str}.xlsx",
              mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          )
        with btn_r3_2:
          pdf_r3 = generate_report_pdf_html(
              pivot_r3,
              subtitle_text=(
                  "भूकरमापक निहाय व दिवसनिहाय प्रलंबित मोजणी प्रकरणे अहवाल"
              ),
          )
          st.download_button(
              label="📄 Download Report 3 (PDF / Print)",
              data=pdf_r3,
              file_name=f"Report_3_Bhukarmapak_{today_file_str}.pdf",
              mime="application/pdf" if WEASYPRINT_AVAILABLE else "text/html",
          )

    with tab5:
      st.markdown("## 📅 रिपोर्ट ४: मोजणी होणेवर शिल्लक अहवाल")
      st.success(f"🎯 **एकूण मोजणी होणेवर शिल्लक प्रकरणे: {len(df_r4)}**")

      if not final_r4_df.empty:
        st.dataframe(final_r4_df, use_container_width=True)

        btn_r4_1, btn_r4_2 = st.columns(2)
        with btn_r4_1:
          excel_r4 = convert_df_to_excel(
              final_r4_df, report_title="मोजणी होणेवर शिल्लक प्रकरणांचा अहवाल"
          )
          st.download_button(
              label="📥 Download Report 4 (Excel)",
              data=excel_r4,
              file_name=f"Report_4_Mojni_Shillak_{today_file_str}.xlsx",
              mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          )
        with btn_r4_2:
          pdf_r4 = generate_report_pdf_html(
              final_r4_df,
              subtitle_text="मोजणी होणेवर शिल्लक प्रकरणांचा अहवाल",
          )
          st.download_button(
              label="📄 Download Report 4 (PDF / Print Layout)",
              data=pdf_r4,
              file_name=f"Report_4_Mojni_Shillak_{today_file_str}.pdf",
              mime="application/pdf" if WEASYPRINT_AVAILABLE else "text/html",
          )

    with tab6:
      st.markdown("## 🚀 रिपोर्ट ५: मागील ३ महिन्यांचा 'क प्रत' प्रोग्रेस अहवाल")
      st.success(
          f"🎯 **मागील ३ महिन्यात पूर्ण झालेली 'क प्रत' प्रकरणे: {len(df_r5)}**"
      )

      if not pivot_r5.empty:
        st.dataframe(pivot_r5, use_container_width=True)

        btn_r5_1, btn_r5_2 = st.columns(2)
        with btn_r5_1:
          excel_r5 = convert_df_to_excel(
              pivot_r5,
              report_title="मागील ३ महिने 'क प्रत' प्रोग्रेस अहवाल",
          )
          st.download_button(
              label="📥 Download Report 5 (Excel)",
              data=excel_r5,
              file_name=f"Report_5_K_Prat_Progress_{today_file_str}.xlsx",
              mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          )
        with btn_r5_2:
          pdf_r5 = generate_report_pdf_html(
              pivot_r5, subtitle_text="मागील ३ महिने 'क प्रत' प्रोग्रेस अहवाल"
          )
          st.download_button(
              label="📄 Download Report 5 (PDF / Print)",
              data=pdf_r5,
              file_name=f"Report_5_K_Prat_Progress_{today_file_str}.pdf",
              mime="application/pdf" if WEASYPRINT_AVAILABLE else "text/html",
          )

    with tab7:
      st.markdown(
          "## ⚠️ रिपोर्ट ६: भूकरमापक प्रलंबित मोजणी १० पेक्षा जास्त प्रकरणे"
          " अहवाल"
      )
      if not df_r6.empty:
        st.warning(
            f"⚠️ **१० पेक्षा जास्त प्रलंबित प्रकरणे असलेल्या भूकरमापकांची संख्या:"
            f" {len(df_r6)}**"
        )
        st.dataframe(df_r6, use_container_width=True)

        btn_r6_1, btn_r6_2 = st.columns(2)
        with btn_r6_1:
          excel_r6 = convert_df_to_excel(
              df_r6,
              report_title=(
                  "भूकरमापक प्रलंबित मोजणी १० पेक्षा जास्त प्रकरणे अहवाल"
              ),
          )
          st.download_button(
              label="📥 Download Report 6 (Excel)",
              data=excel_r6,
              file_name=f"Report_6_Bhukarmapak_10plus_{today_file_str}.xlsx",
              mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          )
        with btn_r6_2:
          pdf_r6 = generate_report_pdf_html(
              df_r6,
              subtitle_text=(
                  "भूकरमापक प्रलंबित मोजणी १० पेक्षा जास्त प्रकरणे अहवाल"
              ),
          )
          st.download_button(
              label="📄 Download Report 6 (PDF / Print)",
              data=pdf_r6,
              file_name=f"Report_6_Bhukarmapak_10plus_{today_file_str}.pdf",
              mime="application/pdf" if WEASYPRINT_AVAILABLE else "text/html",
          )

    # TAB 8: REPORT 7
    with tab8:
      st.markdown(
          "## 📑 रिपोर्ट ७: प्रलंबित प्रकरणांची वैयक्तिक यादी (जिल्हा ➔ तालुका"
          " ➔ भूकरमापक)"
      )
      if not df_r7_display.empty:
        st.info(f"📋 **एकूण वैयक्तिक प्रलंबित प्रकरणे: {len(df_r7_display)}**")
        st.dataframe(df_r7_display, use_container_width=True)

        btn_r7_1, btn_r7_2 = st.columns(2)
        with btn_r7_1:
          excel_r7 = convert_df_to_excel(
              df_r7_display,
              report_title=(
                  "प्रलंबित प्रकरणांची वैयक्तिक यादी (जिल्हा ➔ तालुका ➔"
                  " भूकरमापक) अहवाल"
              ),
          )
          st.download_button(
              label="📥 Download Report 7 (Excel)",
              data=excel_r7,
              file_name=f"Report_7_Vaiktik_Yadi_{today_file_str}.xlsx",
              mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          )
        with btn_r7_2:
          pdf_r7 = generate_report_pdf_html(
              df_r7_display,
              subtitle_text=(
                  "प्रलंबित प्रकरणांची वैयक्तिक यादी (जिल्हा ➔ तालुका ➔"
                  " भूकरमापक) अहवाल"
              ),
          )
          st.download_button(
              label="📄 Download Report 7 (PDF / Print)",
              data=pdf_r7,
              file_name=f"Report_7_Vaiktik_Yadi_{today_file_str}.pdf",
              mime="application/pdf" if WEASYPRINT_AVAILABLE else "text/html",
          )

else:
  st.info("👈 कृपया डाव्या बाजूने Excel फाईल Upload करा.")
